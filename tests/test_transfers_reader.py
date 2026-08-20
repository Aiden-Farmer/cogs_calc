from __future__ import annotations

import pytest

from src.data import FailedRow, Header
from src.transfers.reader import TransferFileReader
from src.transfers.transfer_rows import TransferRow


def _write_workbook(tmp_path, rows, title="Transfers") -> str:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = title
    for row in rows:
        wb.active.append(row)
    path = tmp_path / "transfers.xlsx"
    wb.save(path)
    return str(path)


def _header():
    return Header.transfer_row(
        from_sku=0, to_sku=1, qty=2, date=3, date_format="%Y-%m-%d"
    )


class TestTransferFileReaderInitialization:
    def test_reads_the_named_sheet_not_just_the_active_one(self, tmp_path):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.active.append(["header"])
        other = wb.create_sheet("Transfers")
        other.append(["from", "to", "qty", "date"])
        other.append(["sku-a", "sku-b", 3, "2024-06-01"])
        path = tmp_path / "transfers.xlsx"
        wb.save(path)

        reader = TransferFileReader(
            filename=str(path), sheet_name="Transfers", header=_header()
        )
        rows = list(reader.readline())

        assert len(rows) == 1
        assert isinstance(rows[0], TransferRow)
        assert rows[0].from_sku == "sku-a"

    def test_raises_keyerror_for_unknown_sheet_name(self, tmp_path):
        path = _write_workbook(tmp_path, [["header"], ["sku-a", "sku-b", 3, "2024-06-01"]])

        with pytest.raises(KeyError):
            TransferFileReader(
                filename=path, sheet_name="Nonexistent", header=_header()
            )


class TestTransferFileReaderReadline:
    def test_skips_header_row_and_parses_the_rest(self, tmp_path):
        path = _write_workbook(
            tmp_path,
            [
                ["header"],
                ["sku-a", "sku-b", 3, "2024-06-01"],
                ["sku-c", "sku-d", 5, "2024-06-02"],
            ],
        )

        reader = TransferFileReader(
            filename=path, sheet_name="Transfers", header=_header()
        )
        rows = list(reader.readline())

        assert len(rows) == 2
        assert all(isinstance(r, TransferRow) for r in rows)
        assert [r.from_sku for r in rows] == ["sku-a", "sku-c"]

    def test_bad_row_is_yielded_as_failedrow_not_dropped(self, tmp_path):
        path = _write_workbook(
            tmp_path,
            [
                ["header"],
                ["sku-a", "sku-b", 0, "2024-06-01"],  # qty must be > 0
            ],
        )

        reader = TransferFileReader(
            filename=path, sheet_name="Transfers", header=_header()
        )
        rows = list(reader.readline())

        assert len(rows) == 1
        assert isinstance(rows[0], FailedRow)

    def test_close_frees_the_workbook(self, tmp_path):
        path = _write_workbook(tmp_path, [["header"]])

        reader = TransferFileReader(
            filename=path, sheet_name="Transfers", header=_header()
        )

        reader.close()  # should not raise
